import docx
from docx.shared import Pt
import re

def create_docx(md_file, docx_file):
    doc = docx.Document()
    
    # Simple markdown parser
    with open(md_file, 'r', encoding='utf-8') as f:
        lines = f.readlines()
        
    for line in lines:
        line = line.strip()
        if not line:
            continue
            
        # Headings
        if line.startswith('# '):
            doc.add_heading(line[2:], level=1)
        elif line.startswith('## '):
            doc.add_heading(line[3:], level=2)
        elif line.startswith('### '):
            doc.add_heading(line[4:], level=3)
            
        # Lists
        elif line.startswith('* ') or line.startswith('- '):
            p = doc.add_paragraph(line[2:], style='List Bullet')
            # Handle bold in list items
            bold_pattern = r'\*\*(.*?)\*\*'
            if '**' in line:
                p.clear() # Clear default text, add runs
                parts = re.split(bold_pattern, line[2:])
                for i, part in enumerate(parts):
                    run = p.add_run(part)
                    if i % 2 == 1: # Found bold part
                         run.bold = True
                         
        # Normal Para with Bold handling
        else:
            p = doc.add_paragraph()
            bold_pattern = r'\*\*(.*?)\*\*'
            if '**' in line:
                parts = re.split(bold_pattern, line)
                for i, part in enumerate(parts):
                    run = p.add_run(part)
                    if i % 2 == 1: # Found bold part
                         run.bold = True
            else:
                p.add_run(line)

    doc.save(docx_file)
    print(f"Created {docx_file}")

if __name__ == "__main__":
    create_docx(r"c:\Users\simon\Desktop\ndis case\case_structure_proposal_simon_first.md", 
                r"c:\Users\simon\Desktop\ndis case\case_structure_proposal_simon_first.docx")
