import os, sys, json, hashlib, subprocess, time, email, re
from email import policy
H=os.path.expanduser('~/mnt')
ROOTS=[('CASE',H+'/NDIS Case'),
       ('DESK_FEDPKG',H+'/Desktop/NDIS_FEDERAL_COURT_FINAL_PACKAGE_15_JULY_2026'),
       ('DESK_DRAFTPACK',H+'/Desktop/Reid_ART_draft_pack'),
       ('DESK_ndiscase',H+'/Desktop/ndis case'),
       ('DESK_DUPS',H+'/Desktop/🟫 DUPLICATES'),
       ('DESK_SESSION',H+'/Desktop/SESSION_SUMMARY_16-17_JUL_2026.md')]
OUT=H+'/NDIS Case/Chronology/_corpus'
MAN=H+'/NDIS Case/Chronology/_build/manifest.jsonl'
SKIP_DIRS=['/Chronology/_corpus','/Chronology/_build','/Chronology/_agents','/.git','node_modules','__pycache__']
TEXT={'txt','md','csv','json','mermaid','lst','html','sql','log'}
budget=float(sys.argv[1]) if len(sys.argv)>1 else 150
t0=time.time()
done={}
if os.path.exists(MAN):
    for l in open(MAN,encoding='utf-8'):
        try: r=json.loads(l); done[r['path']]=r
        except: pass
seen_hash={r['sha256']:r['id'] for r in done.values() if r.get('sha256')}
nid=max([r['id'] for r in done.values()]+[0])+1
def files():
    for tag,root in ROOTS:
        if os.path.isfile(root): yield tag,root; continue
        for dp,dn,fn in os.walk(root):
            if any(s in dp for s in SKIP_DIRS): continue
            for f in sorted(fn):
                yield tag,os.path.join(dp,f)
def sha(p):
    h=hashlib.sha256()
    with open(p,'rb') as f:
        for b in iter(lambda:f.read(1<<20),b''): h.update(b)
    return h.hexdigest()
def ext_pdf(p):
    try:
        o=subprocess.run(['pdftotext','-layout',p,'-'],capture_output=True,timeout=90).stdout.decode('utf-8','ignore')
    except Exception as e: return '', 'error:pdftotext'
    try:
        info=subprocess.run(['pdfinfo',p],capture_output=True,timeout=30).stdout.decode('utf-8','ignore')
        pages=int(re.search(r'Pages:\s+(\d+)',info).group(1))
    except: pages=max(1,o.count('\f'))
    st='ok_%dp'%pages
    if len(o.strip())<60*pages: st='scanned_%dp'%pages
    return o, st
def ext_docx(p):
    import docx
    d=docx.Document(p); parts=[x.text for x in d.paragraphs]
    for t in d.tables:
        for r in t.rows: parts.append(' | '.join(c.text for c in r.cells))
    return '\n'.join(parts),'ok'
def ext_eml(p):
    m=email.message_from_binary_file(open(p,'rb'),policy=policy.default)
    hdr='\n'.join(f'{k}: {m.get(k,"")}' for k in ['From','To','Cc','Date','Subject'])
    body=''
    try:
        b=m.get_body(preferencelist=('plain','html'))
        body=b.get_content() if b else ''
    except Exception: pass
    if '<html' in body.lower(): body=re.sub(r'<[^>]+>',' ',body)
    atts=[a.get_filename() for a in m.iter_attachments()]
    return hdr+'\nAttachments: '+', '.join(str(a) for a in atts)+'\n\n'+body,'ok'
def ext_msg(p):
    import extract_msg
    m=extract_msg.Message(p)
    hdr=f'From: {m.sender}\nTo: {m.to}\nCc: {m.cc}\nDate: {m.date}\nSubject: {m.subject}'
    atts=[a.longFilename or a.shortFilename for a in m.attachments]
    b=m.body or ''
    m.close()
    return hdr+'\nAttachments: '+', '.join(str(a) for a in atts)+'\n\n'+b,'ok'
def ext_xlsx(p):
    import openpyxl
    wb=openpyxl.load_workbook(p,read_only=True,data_only=True); out=[]
    for ws in wb.worksheets:
        out.append('## '+ws.title)
        for r in ws.iter_rows(values_only=True):
            if any(c is not None for c in r): out.append(' | '.join('' if c is None else str(c) for c in r))
    return '\n'.join(out),'ok'
mf=open(MAN,'a',encoding='utf-8'); n=0
for tag,p in files():
    if time.time()-t0>budget: print('PAUSED',n); break
    if p in done: continue
    rel=p.replace(H+'/','')
    e=p.rsplit('.',1)[-1].lower() if '.' in os.path.basename(p) else ''
    if os.path.basename(p).startswith('~$'): continue
    try: st_=os.stat(p); h=sha(p)
    except Exception as ex: continue
    r={'id':nid,'path':p,'rel':rel,'root':tag,'ext':e,'size':st_.st_size,'mtime':time.strftime('%Y-%m-%d %H:%M',time.gmtime(st_.st_mtime)),'sha256':h}
    if h in seen_hash:
        r.update(status='duplicate_of',dup_of=seen_hash[h],chars=0)
    else:
        txt,stt='', 'skipped'
        try:
            if e=='pdf': txt,stt=ext_pdf(p)
            elif e=='docx': txt,stt=ext_docx(p)
            elif e=='eml': txt,stt=ext_eml(p)
            elif e=='msg': txt,stt=ext_msg(p)
            elif e=='xlsx': txt,stt=ext_xlsx(p)
            elif e in TEXT and st_.st_size<15_000_000: txt=open(p,encoding='utf-8',errors='ignore').read(); stt='ok'
            elif e in ('png','jpg','jpeg'): stt='image_pending_ocr'
            else: stt='skipped_'+e
        except Exception as ex: stt='error:'+type(ex).__name__
        if txt:
            open(f'{OUT}/{nid:05d}.txt','w',encoding='utf-8').write(f'### SOURCE: {rel}\n### MTIME: {r["mtime"]}\n\n'+txt)
        r.update(status=stt,chars=len(txt))
        seen_hash[h]=nid
    mf.write(json.dumps(r,ensure_ascii=False)+'\n'); mf.flush()
    done[p]=r; nid+=1; n+=1
else:
    print('COMPLETE',n)
